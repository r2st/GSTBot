"""Getting text out of an uploaded file, before anything tries to understand it.

The order is deliberate and it is a cost decision. A PDF exported from Tally
carries its text already; reading it is free, exact, and instant. Only when a
file has no extractable text — a photograph, a scan — does it become worth
spending a vision-model call or a Tesseract pass on it.
"""
from __future__ import annotations

import csv
import io
import logging

logger = logging.getLogger(__name__)

# Anything else is treated as an image and sent to the vision path.
TEXT_CONTENT_TYPES = {"text/plain", "text/csv", "application/csv"}
PDF_CONTENT_TYPES = {"application/pdf"}
EXCEL_CONTENT_TYPES = {
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
IMAGE_CONTENT_TYPES = {"image/jpeg", "image/jpg", "image/png", "image/webp", "image/heic"}

# Below this, a PDF's "text" is page furniture — a header, a page number — and
# the document is really a scan. Measured in characters across the whole file.
_MIN_PDF_TEXT_CHARS = 40

# The most text any upload is allowed to become.
#
# ``MAX_UPLOAD_MB`` bounds the bytes on the wire and does not bound this. A
# spreadsheet is a zip of shared strings, so one cell value repeated across a
# sheet costs almost nothing on disk and a full copy per cell once read: a
# 9 MB .xlsx — comfortably inside the default 15 MB limit — flattens to 149 MB
# of text in 28 seconds and 375 MB of resident memory. A PDF concatenating
# every page does the same more slowly.
#
# That string is not passed through and discarded. It is regex-scanned field by
# field by the heuristic parser, and then written whole into ``invoices.raw_text``
# — so one upload is a worker's memory, a long CPU-bound stretch of it, and a
# row the database has to carry from then on. Uploads run inline in the request
# when Celery is off, which makes it a request that holds all three.
#
# Two million characters is far past any real invoice or purchase register —
# the extraction prompt itself only ever sends the first 12,000 — and short
# enough that the worst case is bounded rather than open. Truncation happens
# while the text is being built, not after, because building it is the cost.
MAX_TEXT_CHARS = 2_000_000


class _TextLimitReached(Exception):
    """Internal: break out of a nested sheet/row walk once the cap is hit."""


def is_image(content_type: str | None, filename: str | None = None) -> bool:
    """Whether this upload should go down the vision/OCR path."""
    if content_type and content_type.lower() in IMAGE_CONTENT_TYPES:
        return True
    if filename:
        lowered = filename.lower()
        return lowered.endswith((".jpg", ".jpeg", ".png", ".webp", ".heic"))
    return False


def from_pdf(content: bytes) -> str:
    """Extract embedded text from a PDF, or ``""`` if it has none."""
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - pypdf is a hard dependency
        logger.warning("pypdf is not installed; cannot read PDF text")
        return ""
    try:
        reader = PdfReader(io.BytesIO(content))
        pages: list[str] = []
        size = 0
        for page in reader.pages:
            pages.append(page.extract_text() or "")
            size += len(pages[-1]) + 1
            if size >= MAX_TEXT_CHARS:
                logger.warning(
                    "PDF text stopped at %s characters (%s pages read)",
                    MAX_TEXT_CHARS,
                    len(pages),
                )
                break
    except Exception as exc:  # noqa: BLE001 - a malformed PDF must not 500
        logger.warning("PDF text extraction failed: %s", exc)
        return ""
    text = "\n".join(pages)[:MAX_TEXT_CHARS].strip()
    return text if len(text) >= _MIN_PDF_TEXT_CHARS else ""


def from_excel(content: bytes) -> str:
    """Flatten a spreadsheet to tab-separated rows.

    Purchase registers arrive as Excel far more often than as anything else,
    and a flattened sheet is something the extraction prompt reads well.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - openpyxl is a hard dependency
        logger.warning("openpyxl is not installed; cannot read spreadsheets")
        return ""
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - a corrupt upload must not 500
        logger.warning("Spreadsheet read failed: %s", exc)
        return ""
    lines: list[str] = []
    size = 0
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                if not any(cell.strip() for cell in cells):
                    continue
                lines.append("\t".join(cells))
                size += len(lines[-1]) + 1
                if size >= MAX_TEXT_CHARS:
                    logger.warning(
                        "Spreadsheet text stopped at %s characters (%s rows read)",
                        MAX_TEXT_CHARS,
                        len(lines),
                    )
                    raise _TextLimitReached
    except _TextLimitReached:
        pass
    finally:
        workbook.close()
    return "\n".join(lines)[:MAX_TEXT_CHARS].strip()


def from_csv(content: bytes) -> str:
    """Normalise a CSV to tab-separated rows."""
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1", errors="replace")
    # A CSV is already as big as it was on the wire, so the cap here is about
    # the row the text ends up in rather than about memory. Applied before the
    # join so the same ceiling holds however the file is shaped.
    rows = list(csv.reader(io.StringIO(decoded[:MAX_TEXT_CHARS])))
    joined = "\n".join(
        "\t".join(row) for row in rows if any(cell.strip() for cell in row)
    )
    return joined[:MAX_TEXT_CHARS].strip()


def ocr_image(content: bytes) -> str:
    """Run Tesseract over an image, or return ``""`` if it is unavailable.

    The fallback for when no vision model is reachable. Both the Python
    binding and the ``tesseract`` binary are optional at runtime: a deployment
    without them still ingests invoices, it just leans entirely on the vision
    model.
    """
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        logger.info("pytesseract/Pillow not installed; skipping OCR fallback")
        return ""
    try:
        with Image.open(io.BytesIO(content)) as image:
            return str(pytesseract.image_to_string(image)).strip()
    except Exception as exc:  # noqa: BLE001 - a missing binary must not 500
        logger.warning("OCR failed: %s", exc)
        return ""


def extract(content: bytes, content_type: str | None, filename: str | None = None) -> str:
    """Best-effort text for an upload of any supported type.

    Returns ``""`` for images and for scans, which is the signal to the parser
    that this document needs the vision path.
    """
    kind = (content_type or "").lower()
    name = (filename or "").lower()

    if kind in PDF_CONTENT_TYPES or name.endswith(".pdf"):
        return from_pdf(content)
    if kind in EXCEL_CONTENT_TYPES or name.endswith((".xlsx", ".xlsm", ".xls")):
        return from_excel(content)
    if kind in {"text/csv", "application/csv"} or name.endswith(".csv"):
        return from_csv(content)
    if kind in TEXT_CONTENT_TYPES or name.endswith(".txt"):
        return content.decode("utf-8", errors="replace")[:MAX_TEXT_CHARS].strip()
    return ""
